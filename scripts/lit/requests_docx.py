"""Write the list of inaccessible full texts (docs/lit-review/fulltext-requests.csv) as a Word document to share.

Usage: uv run --group analysis python scripts/lit/requests_docx.py
Output: docs/lit-review/fulltext-requests.docx
"""

import csv
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

LR = Path(__file__).resolve().parents[2] / "docs" / "lit-review"


def link(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    pr.append(c)
    pr.append(u)
    r.append(pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def main() -> None:
    rows = list(csv.DictReader((LR / "fulltext-requests.csv").open(encoding="utf-8")))
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading("Full texts needed for the literature review", level=1)
    counts = Counter(r["publisher"] for r in rows)
    doc.add_paragraph(
        f"{len(rows)} papers could not be accessed openly ("
        + ", ".join(f"{k} {v}" for k, v in sorted(counts.items(), key=lambda x: -x[1])) + "). "
        "Please save each PDF with the file name in the first column (for example R0041.pdf); "
        "that name links the file to our records.")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, h in zip(table.rows[0].cells, ["Save as", "Publisher", "Year", "Title", "DOI / link"]):
        cell.text = h
    for r in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = r["save_as"], r["publisher"], r["year"], r["title"]
        p = cells[4].paragraphs[0]
        if r["link"]:
            link(p, r["link"], r["doi"])
        else:
            p.text = "no DOI: " + (r["venue"] or "-")
    for row in table.rows:
        for cell, w in zip(row.cells, [Cm(2.2), Cm(2.0), Cm(1.2), Cm(14.0), Cm(6.5)]):
            cell.width = w
    out = LR / "fulltext-requests.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
